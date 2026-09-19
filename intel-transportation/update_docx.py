import docx

doc_path = 'D:/intelligent_transportation/人工智能智慧交通实训报告.docx'
doc = docx.Document(doc_path)

target_idx = -1
for i, p in enumerate(doc.paragraphs):
    if '三、 实训总结与心得体会' in p.text:
        target_idx = i
        break

if target_idx != -1:
    target_p = doc.paragraphs[target_idx]
    
    # 尝试找到之前的二级标题样式
    h2_style = None
    for p in doc.paragraphs:
        if '5. 交通数据时序存储' in p.text:
            h2_style = p.style
            break
            
    h = target_p.insert_paragraph_before('6. 数据库性能调优与离线分析归档（扩展实践）')
    if h2_style:
        h.style = h2_style
        
    target_p.insert_paragraph_before('• TimescaleDB 连续聚合与数据保留：配置了连续聚合（Continuous Aggregates）视图，实现每分钟车流量与平均车速的后台预计算。同时增设数据保留策略，自动清理30天前的原始明细数据，有效降低边缘设备的存储压力，满足企业级系统要求。')
    target_p.insert_paragraph_before('• DuckDB 离线分析与 Parquet 归档：编写并执行了 data_analysis.py 数据分析脚本，引入纯内存的 DuckDB 引擎。通过执行 SQL 对本地数据进行超速违章车辆追踪，并将历史明细零拷贝导出为高压缩比的 Parquet 格式，实现边缘端冷数据离线归档。')
    target_p.insert_paragraph_before('• Pandas 智能数据统计：将 DuckDB 的查询结果对接 Pandas，利用 DataFrame 生成了交通车速的基础数据特征（均值、中位数、极值等）统计报表，完善了“流式在线处理 + 离线归档分析”的完整大数据框架。')
    target_p.insert_paragraph_before('')
    
    doc.save(doc_path)
    print("报告更新成功！")
else:
    print("未找到总结章节，更新失败。")
