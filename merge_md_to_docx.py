import docx
import os
import re

doc_path = 'D:/intelligent_transportation/人工智能智慧交通实训报告_新.docx'
md_path = 'C:/Users/37535/Desktop/智慧交通系统实施与复盘报告.md'

print(f"Reading {md_path}...")
with open(md_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

print(f"Opening {doc_path}...")
doc = docx.Document(doc_path)
doc.add_page_break()
doc.add_heading('四、 智慧交通系统实施与复盘报告', level=1)

def add_md_paragraph(doc, text, list_type=None):
    p = doc.add_paragraph()
    if list_type == 'bullet':
        text = '• ' + text

    # 处理Markdown中的 **粗体**
    parts = text.split('**')
    for i, part in enumerate(parts):
        run = p.add_run(part)
        if i % 2 != 0 and len(parts) > 1:
            run.bold = True
    return p

in_code_block = False
code_content = []

for line in lines:
    stripped = line.strip()
    
    if stripped.startswith('`'):
        if in_code_block:
            # 结束代码块
            add_md_paragraph(doc, ''.join(code_content))
            in_code_block = False
            code_content = []
        else:
            # 开始代码块
            in_code_block = True
        continue
        
    if in_code_block:
        code_content.append(line)
        continue
        
    if not stripped:
        continue
        
    if stripped.startswith('# '):
        doc.add_heading(stripped[2:].replace('**',''), level=2)
    elif stripped.startswith('## '):
        doc.add_heading(stripped[3:].replace('**',''), level=2)
    elif stripped.startswith('### '):
        doc.add_heading(stripped[4:].replace('**',''), level=3)
    elif stripped.startswith('- ') or stripped.startswith('* '):
        add_md_paragraph(doc, stripped[2:], list_type='bullet')
    else:
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            add_md_paragraph(doc, f"{num_match.group(1)}. {num_match.group(2)}")
        elif stripped.startswith('> '):
            add_md_paragraph(doc, stripped[2:])
        else:
            add_md_paragraph(doc, stripped)

doc.save(doc_path)
print("✅ Merge complete! 文件已保存。")
