import sys
from pathlib import Path

from docx import Document


sys.stdout.reconfigure(encoding="utf-8")
report = Path("D:/intelligent_transportation/人工智能智慧交通实训报告_新.docx")
doc = Document(str(report))

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
for index, paragraph in enumerate(doc.paragraphs):
    if 95 <= index <= 110:
        drawings = len(paragraph._p.xpath(".//w:drawing"))
        print(f"{index}|{paragraph.style.name}|drawings={drawings}|{paragraph.text!r}")

required = [
    "8. 重庆绕城高速交通流量预测与大屏可视化（9月9日）",
    "图 8-1 重庆绕城高速交通流量时序预测与可视化监控大屏：",
]
all_text = "\n".join(p.text for p in doc.paragraphs)
for marker in required:
    print(f"marker {marker!r}: {marker in all_text}")
