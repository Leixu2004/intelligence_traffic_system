"""Integrate the storage-platform requirements audit into the main report."""

from __future__ import annotations

import shutil
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
REPORT_PATH = ROOT / "人工智能智慧交通实训报告_新.docx"
BACKUP_PATH = ROOT / "人工智能智慧交通实训报告_新_20260916_更新前备份.docx"
DEPLOYMENT_IMAGE = Path(
    r"C:\Users\37535\xwechat_files\wxid_t46slp02uhwi22_e940\temp\RWTemp\2026-09\be9c956f32b8f5e796e94066718e66a6\f0d447d0ebee62c0e0cb3b5b7e013cec.png"
)
MODULE_IMAGE = Path(
    r"C:\Users\37535\xwechat_files\wxid_t46slp02uhwi22_e940\temp\RWTemp\2026-09\be9c956f32b8f5e796e94066718e66a6\817146984bed2cc02c96a47e8bbdf07c.png"
)

SECTION_TITLE = "数据存储实战课件需求解析与项目符合性评估（9月16日）"
SUMMARY_TITLE = "实训总结与心得体会"


def set_run_font(run, size: float = 10.5, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def replace_paragraph(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run)


def replace_by_prefix(doc: Document, prefix: str, text: str) -> Paragraph:
    paragraph = next((p for p in doc.paragraphs if p.text.strip().startswith(prefix)), None)
    if paragraph is None:
        raise RuntimeError(f"Paragraph not found: {prefix}")
    replace_paragraph(paragraph, text)
    return paragraph


def replace_by_prefixes(doc: Document, prefixes: tuple[str, ...], text: str) -> Paragraph:
    paragraph = next(
        (p for p in doc.paragraphs if any(p.text.strip().startswith(prefix) for prefix in prefixes)),
        None,
    )
    if paragraph is None:
        raise RuntimeError(f"Paragraph not found: {prefixes}")
    replace_paragraph(paragraph, text)
    return paragraph


def insert_before(anchor: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_body(anchor: Paragraph, text: str, *, keep_with_next: bool = False) -> Paragraph:
    paragraph = insert_before(anchor, text, "Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.keep_with_next = keep_with_next
    return paragraph


def add_heading(anchor: Paragraph, text: str, level: int, *, page_break: bool = False) -> Paragraph:
    paragraph = insert_before(anchor, text, f"Heading {level}")
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_caption(anchor: Paragraph, text: str) -> Paragraph:
    paragraph = insert_before(anchor, text, "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = False
    for run in paragraph.runs:
        set_run_font(run, 9)
    return paragraph


def add_picture(
    anchor: Paragraph,
    path: Path,
    width: float,
    *,
    page_break: bool = False,
) -> Paragraph:
    paragraph = insert_before(anchor, style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "D9D9D9")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(
    doc: Document,
    anchor: Paragraph,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 8.6,
    center_columns: set[int] | None = None,
) -> None:
    center_columns = center_columns or set()
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, (cell, text, width) in enumerate(zip(header_row.cells, headers, widths)):
        set_cell_width(cell, width)
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_run_font(run, font_size, bold=True, color="FFFFFF")

    for row_index, values in enumerate(rows):
        row_cells = table.add_row().cells
        fill = "F2F6FA" if row_index % 2 else "FFFFFF"
        for column_index, (cell, text, width) in enumerate(zip(row_cells, values, widths)):
            set_cell_width(cell, width)
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(text)
            set_run_font(run, font_size)

    anchor._p.addprevious(table._tbl)
    spacer = insert_before(anchor, style="Normal")
    spacer.paragraph_format.space_after = Pt(6)


def remove_existing_section(doc: Document, anchor: Paragraph) -> None:
    start = next((p for p in doc.paragraphs if p.text.strip() == SECTION_TITLE), None)
    if start is None:
        return
    element = start._element
    while element is not None and element is not anchor._element:
        next_element = element.getnext()
        element.getparent().remove(element)
        element = next_element


def normalize_heading_styles(doc: Document) -> None:
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def normalize_heading_labels(doc: Document) -> None:
    """Let the document's multilevel list own heading numbers."""
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name
        text = paragraph.text.strip()
        if not text:
            continue
        if style_name == "Heading 1":
            clean = re.sub(r"^[一二三四五六七八九十]+、\s*", "", text)
        elif style_name == "Heading 2":
            clean = re.sub(r"^\d+\.\s*", "", text)
        elif style_name == "Heading 3":
            clean = re.sub(r"^\d+(?:\.\d+)+\s*", "", text)
        else:
            continue
        if clean != text:
            replace_paragraph(paragraph, clean)


def main() -> None:
    for path in (REPORT_PATH, DEPLOYMENT_IMAGE, MODULE_IMAGE):
        if not path.exists():
            raise FileNotFoundError(path)

    if not BACKUP_PATH.exists():
        shutil.copy2(REPORT_PATH, BACKUP_PATH)

    doc = Document(str(REPORT_PATH))
    normalize_heading_styles(doc)
    normalize_heading_labels(doc)

    replace_by_prefix(
        doc,
        "本实训报告基于",
        "本实训报告基于“人工智能智慧交通”项目，记录从视频采集、车辆与车牌识别、违章事件生成，到Kafka异步流转、TimescaleDB时序存储、PyTorch时序预测、ONNX服务化和Streamlit WebGIS展示的工程实践。当前版本已形成OpenCV/YOLO/PaddleOCR—Kafka—TimescaleDB—FastAPI/ONNX—Streamlit的单机原型链路，并使用Pandas与DuckDB完成离线分析。中国交通流LSTM模型属于研究候选，边缘侧已完成通用ONNX导出与CPU验证，尚未形成目标设备量化制品。课件提出的Flink SQL实时计算、Milvus Lite法规知识库、Traffic Cop Agent、CrewAI与Qwen-VL应急协同仍需建设，本报告据此给出满足度审计和分阶段实施计划。",
    )
    replace_by_prefixes(
        doc,
        ("• 服务一键部署", "• 服务编排与当前边界"),
        "• 服务编排与当前边界：项目主docker-compose.yml已编排TimescaleDB、Kafka、Python数据消费者、预测API和Streamlit大屏，可完成现有原型的一键启动。课件要求的Flink Standalone、Flink SQL和Traffic Cop Agent尚未进入主Compose；Milvus Lite应作为Agent进程内嵌组件使用持久化卷，而不是独立集群。",
    )
    replace_by_prefix(
        doc,
        "• PaddleOCR框架入门",
        "• PaddleOCR框架入门：学习并应用PaddleOCR检测与识别流水线，当前项目主路径使用PP-OCRv6模型资产，完成车牌区域检测、透视校正、图像增强、字符识别与格式校验。",
    )
    replace_by_prefixes(
        doc,
        ("• 模型转换与量化", "• 模型导出与运行时验证"),
        "• 模型导出与运行时验证：已将YOLOv8车辆检测和车牌检测模型导出为ONNX，并完成ONNX checker、ONNX Runtime CPU加载和数值对比。OCR仍使用PaddleOCR/PaddleX运行时；仓库中的FP16、INT8和TensorRT入口属于后续设备派生制品流程，现有ONNX文件不应表述为已量化模型。",
    )
    replace_by_prefixes(
        doc,
        ("• 边缘适配实战", "• 边缘适配准备"),
        "• 边缘适配准备：项目已建立Jetson TensorRT、FP16、INT8和Atlas等设备路线的制品契约、导出入口与性能验收规范。由于尚未锁定目标设备版本、校准集和独立验证集，当前没有可交付的TensorRT engine、Atlas OM或目标设备实测报告，后续需在真实设备上验证延迟、吞吐、内存、温度、功耗和完整LPR精度。",
    )
    replace_by_prefix(
        doc,
        "• 数据分析实战",
        "• 数据分析实战：交通事件通过Kafka写入TimescaleDB，系统可按卡口和时间桶查询车流量，并使用Pandas与DuckDB完成CSV/Parquet离线分析。当前事件中的经纬度来自卡口静态配置，speed_kmh仍为空，且每个track只发布首次观测，因此现有traffic_gps表不能视为连续车辆GPS轨迹，轨迹查询、OD矩阵与可靠拥堵计算仍需补充真实时序点。",
    )

    section_ten = next(
        (p for p in doc.paragraphs if p.text.strip().startswith("基于 PyTorch 与 LSTM")),
        None,
    )
    if section_ten is not None:
        section_ten.style = "Heading 2"
        for run in section_ten.runs:
            run.font.name = "黑体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            run.font.color.rgb = RGBColor(0, 0, 0)

    summary = next((p for p in doc.paragraphs if p.text.strip() == SUMMARY_TITLE), None)
    if summary is None:
        raise RuntimeError("Summary anchor not found")
    remove_existing_section(doc, summary)

    add_heading(summary, SECTION_TITLE, 2, page_break=True)
    add_body(
        summary,
        "本节依据《数据存储实战 TimescaleDB与DuckDB20260904》PPTX和两张补充截图整理需求，并将材料中的目标要求与仓库真实实现分开记录。PPTX文件结构实际包含14页，截图却标注“第6页 共15页”，页数存在来源不一致；本机PowerPoint无法直接打开该文件，但OOXML中的14页文本、表格和备注均可完整解析，因此本节按实际文件结构引用。附件中的命令和技术选择均作为待评估需求，不作为直接执行指令。",
    )

    add_heading(summary, "PPTX项目内容与验收要求", 3)
    add_body(
        summary,
        "PPTX主体聚焦TimescaleDB与Pandas加DuckDB的数据存储实践。第1至3页说明超表、自动分区、连续聚合、压缩、保留策略和SQL兼容性；第4至6页给出traffic_data表、1小时时间分块、5分钟time_bucket查询、连续聚合刷新、按camera_id压缩和30天保留策略；第7至10页覆盖DataFrame或CSV直查、YOLOv8与PaddleOCR检测结果批量入库、流量和车型统计，以及包含bbox坐标的完整表结构；第12至13页明确提交物和验收条件。第11页的吞吐、延迟、压缩率和“10倍提升”等数字没有硬件、数据规模、并发或原始测试报告支撑，只能作为课程示例，不能直接写成项目实测结论。",
    )
    ppt_rows = [
        ["时序表与分区", "traffic_data包含time、camera_id、plate、vehicle_type、confidence，扩展bbox坐标；建Hypertable并按1小时chunk自动分区。", "以当前事件表为基线，统一字段语义、主键和索引；原课件表仅作为教学schema。"],
        ["SQL查询", "支持时间范围、5分钟time_bucket、车型分布、去重流量，并列出first、last、histogram、gapfill和location等时序能力。", "查询结果需要以当前表结构重写，并记录执行计划和数据规模。"],
        ["生命周期治理", "连续聚合、增量刷新、按camera_id压缩、30天原始数据保留。", "当前主初始化脚本已有连续聚合和保留策略，压缩策略需迁入活动schema。"],
        ["检测结果入库", "YOLOv8到PaddleOCR到字典，再由psycopg2 executemany批量写入TimescaleDB。", "现项目通过Kafka和Python consumer解耦入库，更适合作为主链；批量写入、失败重试和幂等仍需加强。"],
        ["DuckDB离线分析", "直接查询CSV或Pandas DataFrame，完成历史统计并导出CSV或Excel，可扩展Parquet归档。", "已有脚本和产物，但路径、数据库名和表名需与当前项目统一，并纳入可重复任务。"],
        ["提交物与硬验收", "3个SQL脚本、带连接池和重试的写入脚本、存储设计文档、DuckDB脚本；1000条以上写入、成功率100%、time_bucket查询小于50毫秒、统计含去重、文档含ER图。", "当前缺少同一环境下的完整基准报告。小于50毫秒须绑定数据量、索引、并发、硬件和冷热缓存条件后再验收。"],
    ]
    add_table(
        doc,
        summary,
        ["主题", "课件明确要求", "工程解释与约束"],
        ppt_rows,
        [1.05, 2.9, 2.2],
        font_size=8.5,
        center_columns={0},
    )

    add_heading(summary, "补充截图中的平台目标", 3)
    add_body(
        summary,
        "第一张截图要求提供docker-compose.yml，并以docker-compose up -d在5分钟内获得PostgreSQL、TimescaleDB、Milvus Lite和Flink单机版环境。Flink采用Standalone单机伪分布式，支持Checkpoint与窗口计算，不依赖K8s或Yarn；开发栈限定为Python和SQL，实时计算使用Flink SQL。Milvus Lite在另一张图中又被定义为嵌入式且无需独立集群，因此实施时应把它放在Python Agent进程内并挂载持久化目录，而不是创建独立Milvus服务。",
    )
    add_picture(summary, DEPLOYMENT_IMAGE, 6.15)
    add_caption(summary, "图 11-1 环境部署方式与Flink技术约束（用户提供截图）")
    add_body(
        summary,
        "第二张截图跨越两个页面，上半部的“数据采”和下半部的“集与接入”属于同一行，即“数据采集与接入”。完整目标包括视频流、GPS轨迹和卡口数据多源接入，TimescaleDB时序存储，Flink SQL车速统计、拥堵检测、分级预警和CEP，Pandas加DuckDB离线特征工程，未来1至4小时LSTM或Transformer预测，LangChain加Milvus Lite法规RAG，CrewAI加Qwen-VL事故图文研判，以及Streamlit或ECharts WebGIS可视化。核心Traffic Cop Agent需完成“理解、检索、预测、决策”闭环，并通过只读SQL查询器、预测模型和法规检索器生成带依据的结构化疏导方案。",
    )
    add_picture(summary, MODULE_IMAGE, 4.25, page_break=True)
    add_caption(summary, "图 11-2 功能模块、Traffic Cop Agent与知识库目标（用户提供跨页截图）")

    add_heading(summary, "当前项目满足度", 3, page_break=True)
    add_body(
        summary,
        "截至2026年9月16日，项目已经形成单机可运行原型，但尚未达到截图描述的完整平台。现有主链为OpenCV与YOLO/PaddleOCR感知、Kafka事件流、Python consumer、TimescaleDB、FastAPI/ONNX和Streamlit WebGIS。下面的结论以仓库代码、配置、测试和当前部署条件为准。",
    )
    compliance_rows = [
        ["数据采集与多源接入", "视频、GPS轨迹、卡口多源，经Kafka接入", "单路VideoCapture、YOLO跟踪和两类Kafka事件可用；配置一次只运行一个源。经纬度是卡口静态值，速度为空，每个track只发布首次观测。", "部分满足。需SourceAdapter、多路并发、断流重连、真实GPS/卡口协议和连续轨迹事件。"],
        ["TimescaleDB存储", "轨迹查询、特征管理、SQL分区和自动聚合", "Compose固定TimescaleDB PG16，已有两张Hypertable、索引、1分钟连续聚合、刷新与保留策略。", "大体满足原型。活动初始化脚本缺压缩，尚无独立特征表；traffic_gps名称与数据语义不符。"],
        ["实时计算与预警", "Flink Standalone、Flink SQL、Checkpoint、窗口和CEP", "当前由Python consumer写明细、Timescale连续聚合做1分钟统计；主Compose和依赖中没有Flink。", "不满足。没有可靠车速、拥堵状态机、分级告警和Checkpoint恢复。"],
        ["DuckDB离线工程", "历史流量、OD矩阵、时序特征，单机分析", "已有DuckDB、Pandas、CSV SQL、Parquet归档和描述统计。部分脚本硬编码Windows路径，旧表名和数据库名与现状不一致。", "部分满足。需统一schema、环境变量和可调度任务；OD矩阵仍缺真实轨迹。"],
        ["AI智能预测", "未来1至4小时流量和拥堵传播，ONNX部署", "PyTorch训练、独立评估、ONNX Runtime和FastAPI链较完整。默认中国候选模型只预测下一个20分钟桶，大屏递归4步约80分钟。", "部分满足。模型是研究候选，缺授权现场数据、验证过的12步或16步预测和空间传播特征。"],
        ["Traffic Cop Agent", "自然语言查询、法规RAG、预测和结构化决策", "代码、配置和依赖中没有LangChain、Milvus Lite或FAISS，也没有Agent API、法规库、只读SQL工具和引用链。", "不满足。应先实现可审计的只读工具链，再开放方案生成。"],
        ["应急处置", "Qwen-VL图文分析、CrewAI多Agent、多部门协同", "旧演示页面只有静态告警表和按钮，没有事故理解、协作工作流、审批或执行回执。", "不满足。真实信号控制和分流建议必须设置人工审批和权限边界。"],
        ["可视化交互", "实时热力图、Agent对话、大屏、WebGIS", "Streamlit主大屏已有KPI、趋势、预测、车型、路口热力矩阵、高德WebGIS、识别记录和证据图。", "部分满足且MVP较完整。缺Agent对话、事故时间线、告警闭环；当前图表为Plotly而非ECharts。"],
        ["Compose一键部署", "5分钟获得完整数据环境", "主Compose静态校验通过，含TimescaleDB、Kafka、consumer、prediction-api和dashboard；Docker daemon当前未运行。", "部分满足。缺Flink和Agent进程，尚未执行冷启动、健康收敛和端到端5分钟验收。"],
    ]
    add_table(
        doc,
        summary,
        ["模块", "材料要求", "当前事实", "结论与缺口"],
        compliance_rows,
        [1.0, 1.45, 2.05, 1.65],
        font_size=8.0,
        center_columns={0},
    )

    add_heading(summary, "核心质量风险", 3)
    add_body(
        summary,
        "• 数据语义风险：traffic_gps当前保存的是卡口固定坐标和单次车辆观测，不是连续GPS轨迹；分钟车流量依赖tracker ID生命周期，断轨或重识别会改变计数，speed_kmh为空也使平均速度和拥堵判断失去可靠输入。",
    )
    add_body(
        summary,
        "• 识别与规则风险：在线车牌识别只在压线或违停触发后执行，尚不是常态过车LPR。压线规则只判断包围框底边是否进入固定像素带，没有前后位置、方向和冷却状态；违停使用系统时钟，不适合离线视频倍速，也缺少离场清理。",
    )
    add_body(
        summary,
        "• 模型与数据风险：中国LSTM使用公开匿名收费站数据训练，受非商业研究许可和现场泛化限制；单步结果递归外推会累积误差，不能直接宣称已达到未来1至4小时预测。正式上线需要授权现场长周期数据、滚动回测、漂移监测和可回滚版本。",
    )
    add_body(
        summary,
        "• 部署与安全风险：当前默认数据库凭据和API鉴权仅适合本地实训，CORS范围较宽；Python consumer缺少DLQ、批量写入治理和完整可观察性。Docker daemon未启动，因此本次只能确认Compose静态有效，不能确认容器冷启动和故障恢复。",
    )
    add_body(
        summary,
        "• 验证边界：在项目环境中，预测模块34项、Dashboard 7项、Kafka与事件契约4项测试通过，两份Compose静态校验通过。另一次使用缺少psycopg2、fastapi和onnx的解释器运行完整测试时出现依赖加载失败，说明验收命令和锁定环境必须一并交付。当前没有Kafka到Flink到TimescaleDB的容器级集成测试、真实视频精度集和实时告警SLA测试。",
    )

    add_heading(summary, "分阶段实现计划", 3)
    add_body(
        summary,
        "实施顺序先解决数据语义和可重复部署，再增加Flink、预测、Agent和应急协同。工期为单人开发估算，不包含现场数据获取、目标设备采购、第三方模型审批和生产系统联调时间。",
    )
    roadmap_rows = [
        ["P0", "部署与数据基线", "统一track_point、vehicle_passage、violation_event、alert_event语义；加入source_id、event_id、事件时间和幂等键。把压缩策略迁入活动SQL，修正DuckDB脚本和旧表名；Compose增加profiles、健康检查、环境模板和smoke test。", "Compose配置、SQL幂等迁移、1000条批写、失败重试、去重和查询基准。", "1至2天"],
        ["P1", "Flink SQL实时层", "增加JobManager、TaskManager和SQL初始化；配置Kafka与JDBC connector、Checkpoint持久卷、事件时间和watermark。实现1分钟与5分钟TUMBLE/HOP、持续低速CEP或MATCH_RECOGNIZE，并输出traffic_metrics和traffic_alerts。", "Kafka到Flink到TimescaleDB集成测试、Checkpoint恢复、迟到数据和告警升级测试。", "3至5天"],
        ["P2", "多源采集与1至4小时预测", "以SourceAdapter接入RTSP、本地视频、摄像头、GPS和卡口协议；补测速标定与检测线过车口径。以20分钟12步或15分钟16步训练LSTM/Transformer，引入天气、节假日和邻接路段特征。", "断流重连、源级隔离、rolling backtest、基线对比、逐预测步指标和ONNX一致性。", "4至7天"],
        ["P3", "Traffic Cop RAG与应急研判", "新增LangChain Agent API，进程内使用Milvus Lite持久化法规版本、来源和分块；只开放只读Timescale SQL、预测API和法规检索工具。接入Qwen-VL做事故图像结构化抽取，再以CrewAI拆分研判、方案和校核角色。", "回答必须带法规引用和工具审计；危险SQL拒绝；信号灯、分流和限行建议必须人工审批。", "5至10天"],
        ["P4", "可视化与交付验收", "Streamlit增加Agent对话、事故时间线、告警确认与解除、审批和执行回执；按验收需要继续使用Plotly加高德WebGIS，或补充ECharts组件。统一日志、指标、鉴权、密钥和资源限制。", "5分钟一键启动、4小时预测、RAG引用、应急审批、全链路审计和故障降级演练。", "3至5天"],
    ]
    add_table(
        doc,
        summary,
        ["阶段", "目标", "主要交付", "验收重点", "估算"],
        roadmap_rows,
        [0.55, 1.05, 2.55, 1.55, 0.55],
        font_size=7.8,
        center_columns={0, 4},
    )

    add_heading(summary, "完成定义", 3)
    add_body(
        summary,
        "完整平台达标条件如下：预拉取镜像后执行docker compose up -d，5分钟内所有健康检查通过；视频、GPS和卡口事件按统一版本契约进入Kafka；Flink SQL可从Checkpoint恢复并正确输出窗口指标和分级告警；TimescaleDB完成分区、聚合、压缩、保留和可复现基准；DuckDB使用当前schema生成可核验结果；授权现场模型经独立测试和滚动回测稳定优于基线并覆盖1至4小时；Agent保留法规来源、工具审计和人工审批记录；WebGIS展示指标、预测、事故、告警生命周期和执行回执。",
    )

    replace_by_prefixes(
        doc,
        ("经过为期多天的实训", "经过本阶段实训"),
        "经过本阶段实训，我已将计算机视觉、车牌识别、时序预测与Kafka、TimescaleDB、DuckDB、FastAPI和Streamlit等工程能力连接成可运行的单机原型。实践同时表明，模型文件、接口可调用和页面可展示并不等同于生产系统已经完成：真实GPS轨迹、可靠测速与过车口径、Flink SQL事件时间计算、长时预测、法规RAG和应急审批闭环都需要独立的数据契约、测试和验收证据。后续将按照第11节的P0至P4路线推进，优先固化数据语义和部署基线，再建设实时计算、现场模型、Traffic Cop Agent与可审计的应急协同能力。",
    )

    doc.save(str(REPORT_PATH))
    print(REPORT_PATH)
    print(BACKUP_PATH)


if __name__ == "__main__":
    main()
