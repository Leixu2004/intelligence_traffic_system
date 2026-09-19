"""Append the 2026-09-10 dashboard integration record to the training report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
REPORT_PATH = ROOT / "人工智能智慧交通实训报告_新.docx"
RUNTIME_IMAGE = ROOT / "assets" / "dashboard" / "20260910_streamlit_dashboard_runtime.png"
DATAFLOW_IMAGE = ROOT / "assets" / "dashboard" / "20260910_streamlit_dashboard_dataflow.png"


def insert_before(anchor: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    return paragraph


def add_body(anchor: Paragraph, text: str) -> Paragraph:
    return insert_before(anchor, text, "Normal")


def add_caption(anchor: Paragraph, text: str) -> Paragraph:
    paragraph = add_body(anchor, text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.size = Inches(0.13)
    return paragraph


def add_picture(anchor: Paragraph, path: Path, width: float) -> Paragraph:
    paragraph = insert_before(anchor, style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    return paragraph


def main() -> None:
    if not REPORT_PATH.exists():
        raise FileNotFoundError(REPORT_PATH)
    if not RUNTIME_IMAGE.exists() or not DATAFLOW_IMAGE.exists():
        raise FileNotFoundError("Report images have not been generated")

    doc = Document(str(REPORT_PATH))
    existing = [p for p in doc.paragraphs if "Streamlit 与 WebGIS 可视化大屏集成（9月10日）" in p.text]
    if existing:
        raise RuntimeError("The 2026-09-10 dashboard section already exists")

    anchor = next(
        p for p in doc.paragraphs if p.text.strip() == "三、 实训总结与心得体会"
    )

    insert_before(anchor, "9. Streamlit 与 WebGIS 可视化大屏集成（9月10日）", "Heading 2")
    add_body(
        anchor,
        "• 任务目标与资料分析：结合《可视化大屏 Streamlit 与 WebGIS》PPT 以及 Streamlit、WebGIS 和交通流量大屏配套文档，明确本次任务需要完成宽屏监测界面、四项核心指标、交通趋势图、车辆构成分析、卡口地图和定时刷新能力。文档中提到的 Vue、Flink、InfluxDB、Docker Compose 等属于大型系统扩展方案，本次按照当前项目已有的 FastAPI、TimescaleDB、DuckDB 和 Python 运行基础，优先落地可运行的 Streamlit MVP。",
    )
    add_body(
        anchor,
        "• 展示层接入规划：在项目根目录新增 dashboard 模块，将 Streamlit 作为独立的只读展示进程。现有 YOLOv8 车辆检测、PaddleOCR 车牌识别、Kafka 消费和 TimescaleDB 写入链路保持原样运行，大屏通过已有 FastAPI 的 /api/traffic_trend 接口读取聚合数据，从而避免界面刷新阻塞感知和入库任务。",
    )
    add_body(
        anchor,
        "• 数据加载与降级处理：编写 dashboard/data_loader.py 统一不同来源的数据字段和时间格式，并使用 @st.cache_data(ttl=60) 缓存查询结果。数据优先读取 FastAPI/TimescaleDB，接口不可用时依次回退到本地 vehicle_records.csv、traffic_archive.parquet 和固定随机种子的内置示例数据；因此在数据库、Kafka 或网络地图服务未启动的冷启动环境中，页面仍能正常展示并给出数据来源提示。",
    )
    add_picture(anchor, RUNTIME_IMAGE, 6.25)
    add_caption(anchor, "图 9-1 2026年9月10日 Streamlit 交通流量监测大屏运行效果（冷启动示例数据验收）")
    add_body(
        anchor,
        "• 页面功能实现：在 dashboard/app.py 中完成暗色宽屏布局、监测卡口筛选、自动刷新间隔设置和缓存清理操作。页面顶部展示监测车辆数、平均速度、高流量卡口和异常/告警四项 KPI；中部使用 Plotly 绘制车流历史趋势和短时预测曲线，并以环形图展示车辆构成；底部展示 WebGIS 卡口态势、最新车牌识别记录和违章取证图片。另增加 pages/1_交通流量大屏.py 作为 Streamlit 多页面入口，便于后续扩展流量预测和违章检索页面。",
    )
    add_body(
        anchor,
        "• WebGIS 实现方式：编写 dashboard/map_component.html 和 map_component.py，将高德地图 JS API 2.0 的 Key、安全密钥、卡口经纬度和统计指标转换为前端 JSON，支持暗色底图、卡口 Marker、点击信息窗体和热力图层。高德配置通过 AMAP_KEY 与 AMAP_SECURITY_CODE 环境变量注入；未配置 Key 时自动切换到 Streamlit 本地地图，以保证开发和验收阶段不依赖第三方鉴权。",
    )
    add_picture(anchor, DATAFLOW_IMAGE, 6.25)
    add_caption(anchor, "图 9-2 可视化展示层数据链路与外部服务不可用时的降级路径")
    add_body(
        anchor,
        "• 运行验证与结果：在项目专用 .venv-dashboard 环境中安装 Streamlit、Plotly、Pandas、Requests、DuckDB、psycopg2-binary、Matplotlib 等依赖，执行 streamlit run dashboard\\app.py --server.port 8501 启动大屏。通过 compileall 编译检查、2 项数据层单元测试、/_stcore/health 健康检查以及浏览器首页和多页面实测，确认 KPI、趋势图、车辆构成、地图兜底、抓拍表和取证图片均能渲染。实际验收地址为 http://127.0.0.1:8501。",
    )
    add_body(
        anchor,
        "• 工程化成果：本次完成了从现有交通识别和时序存储能力到交互式展示层的第一次完整打通，形成了可独立启动、可离线演示、可接入实时 API、可替换高德地图配置的可视化模块。后续生产化工作可以在此基础上继续补充 API 鉴权、数据库连接池、真实高德域名白名单、时间范围筛选以及更细粒度的预测服务接口。",
    )

    doc.save(str(REPORT_PATH))
    print(REPORT_PATH)


if __name__ == "__main__":
    main()

