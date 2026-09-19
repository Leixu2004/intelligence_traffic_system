from pptx import Presentation
import sys

def read_pptx(file_path):
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text)
        print("\n")

if __name__ == '__main__':
    read_pptx(r'C:\Users\37535\xwechat_files\wxid_t46slp02uhwi22_e940\msg\file\2026-09\数据存储实战_TimescaleDB与DuckDB20260904.pptx')
