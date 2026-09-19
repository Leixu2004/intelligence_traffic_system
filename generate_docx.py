import docx

doc = docx.Document('D:/实训报告以及相关/实训报告.docx')

def set_p_text(p_idx, text):
    p = doc.paragraphs[p_idx]
    if len(p.runs) > 0:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.text = text

set_p_text(18, '专 业 名 称  ： [在此填写专业]')
set_p_text(19, '课 程 名 称  ： 人工智能智慧交通项目实训')
set_p_text(20, '指 导 教 师  ： 赵国权')
set_p_text(22, '学 生 学 号  ： [在此填写学号]')
set_p_text(23, '学 生 姓 名  ： [在此填写姓名]')
set_p_text(26, '二○二六年九月')
set_p_text(38, '本人签名：                      日期：       2026.09.04')

abstract_text = '本实训报告基于“人工智能智慧交通”项目，详细记录了从基础开发环境搭建、AI视觉模型开发到边缘端部署及数据存储优化的全过程。项目前期利用Docker与Docker Compose实现了一键式环境部署，并结合OpenCV完成了基础视频流处理；核心部分深入应用了YOLOv8目标检测框架与PaddleOCR技术，通过交通数据集的标注、微调训练，成功构建了高精度的“车辆/行人检测”与“车牌识别抓拍”端到端流水线。针对智慧交通对高实时性的要求，项目引入ONNX Runtime对模型进行量化加速，完成了轻量级的边缘计算盒部署适配；最后，结合TimescaleDB时序数据库与分析引擎，实现了海量交通识别数据的超表存储及高性能SQL时序查询，完成交通流量的时序统计。本次实训打通了“云-边-端”的智慧交通技术链路，显著提升了综合工程实践能力。'
keywords_text = '关键词：智慧交通；YOLOv8目标检测；PaddleOCR车牌识别；ONNX边缘部署；TimescaleDB'

set_p_text(41, abstract_text)
set_p_text(42, keywords_text)

for i in range(43, len(doc.paragraphs)):
    set_p_text(i, '')

doc.add_page_break()
doc.add_heading('一、 实习/实训目的与背景', level=1)
doc.add_paragraph('随着智慧城市的发展，人工智能在交通领域的应用愈发重要。本次“人工智能智慧交通”项目实训旨在通过真实场景的工程实践，掌握从云端环境配置、深度学习模型（目标检测与OCR）训练与集成、边缘端轻量化部署，到后端时序数据库存储与高性能数据分析的全链路技术。')

doc.add_heading('二、 实训内容与项目实践', level=1)
doc.add_paragraph('结合每日实训任务，我在本项目中主要完成了以下五个阶段的核心工作：')

doc.add_heading('1. 实训环境搭建与系统基础部署（8月31日）', level=2)
doc.add_paragraph('• 开发环境配置：在本地环境中安装并配置了VS Code及Docker，并在VS Code中安装配置了PostgreSQL与Kubernetes相关开发插件，为后续代码编写与服务调试提供支持。')
doc.add_paragraph('• 服务一键部署：编写并执行了 docker-compose up -d 命令，一键启动了项目所需的基础组件栈，包括PostgreSQL关系型数据库、TimescaleDB时序数据库、Milvus Lite向量数据库以及Flink流处理引擎的单机版环境。')
doc.add_paragraph('• 视频流捕获：学习OpenCV基础语法，编写脚本实现了交通监控视频流的捕获与帧处理，为后续的AI模型推理提供了原始图像数据输入。')

doc.add_heading('2. 计算机视觉基础与目标检测模型开发（9月1日）', level=2)
doc.add_paragraph('• YOLOv8框架研究：深入解析了YOLOv8模型的核心架构（包含Backbone主干网络、Neck特征融合层以及Head检测头），并成功调用了官方预训练模型进行初步测试。')
doc.add_paragraph('• 模型微调与实战：针对真实交通场景，完成了交通数据集（车辆、行人等目标）的标注工作。基于标注数据对YOLOv8进行了微调（Fine-tuning）训练，成功开发并集成了一个具备高鲁棒性的车辆与行人检测模型，能够在复杂路况下准确框选目标。')

doc.add_heading('3. OCR技术应用与车牌违章抓拍集成（9月2日）', level=2)
doc.add_paragraph('• PaddleOCR框架入门：学习并掌握了业界领先的PaddleOCR框架，特别针对最新的PP-OCRv5架构进行了深入剖析，理清了“文本检测 + 文本识别”的流水线机制。')
doc.add_paragraph('• 车牌识别系统实战：将YOLOv8车辆检测模型与PaddleOCR相结合，开发了车牌自动识别系统。实现了从图像中定位车辆、裁剪车牌区域，到最终精准输出车牌字符的完整链路，为违章抓拍业务提供了核心算法支持。')

doc.add_heading('4. 边缘部署优化与模型加速（9月3日）', level=2)
doc.add_paragraph('• 模型转换与量化：针对交通路口边缘设备算力有限的问题，引入了ONNX Runtime技术。将基于PyTorch训练出的YOLOv8和OCR模型成功转换为跨平台的ONNX格式，并实施了模型量化操作以缩减模型体积。')
doc.add_paragraph('• 边缘适配实战：完成了模型的轻量级边缘部署工作，对Nvidia Jetson / Atlas等边缘计算盒子进行了环境适配与性能调优，大幅提升了模型在边缘侧的推理帧率（FPS）与响应速度。')

doc.add_heading('5. 交通数据时序存储与高性能分析（9月4日）', level=2)
doc.add_paragraph('• 时序数据库构建：学习并引入了专为时间序列设计的TimescaleDB。在数据库中创建了用于存储车辆轨迹与识别记录的超表（Hypertable），并配置了基于时间维度的自动分区策略。')
doc.add_paragraph('• 数据分析实战：将前端AI视觉模型识别产生的数据（如时间、车牌号、GPS坐标、速度等）通过代码实时写入TimescaleDB。结合Pandas与DuckDB单机高性能分析引擎，编写SQL语句实现了交通流量的时序查询、拥堵状态统计及可视化数据的输出。')

doc.add_heading('三、 实训总结与心得体会', level=1)
doc.add_paragraph('经过为期多天的实训，我成功将零散的人工智能知识（计算机视觉、深度学习）与工程化技术（Docker容器化、ONNX边缘部署、TimescaleDB数据存储）进行了高度融合。\n\n在实践过程中，我不仅掌握了YOLOv8与PaddleOCR在真实交通抓拍场景中的应用技巧，更深刻理解了AI模型从“实验室训练”走向“工业级边缘部署”所需经历的量化与加速过程。最后的数据存储实战，让我认识到在面对海量物联网与交通流数据时，选择合适的时序数据库与分析引擎（TimescaleDB + DuckDB）对系统整体性能的巨大提升。本次实训极大提升了我的代码落地能力和系统架构思维。')

doc.save('D:/intelligent_transportation/人工智能智慧交通实训报告.docx')
