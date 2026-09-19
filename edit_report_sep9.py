from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


REPORT = Path(r"D:\intelligent_transportation\人工智能智慧交通实训报告_新.docx")
IMAGE = Path(
    r"C:\Users\37535\Desktop\实训留痕迹\基本直线绕城高速重庆\traffic_prediction_dashboard.png"
)


def insert_before(reference, paragraph):
    reference._p.addprevious(paragraph._p)


def add_text_before(doc, reference, text, style="Normal"):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    insert_before(reference, paragraph)
    return paragraph


def add_image_before(doc, reference, image_path, width_inches=5.8):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    insert_before(reference, paragraph)
    return paragraph


def main():
    if not REPORT.exists():
        raise FileNotFoundError(REPORT)
    if not IMAGE.exists():
        raise FileNotFoundError(IMAGE)

    doc = Document(str(REPORT))
    existing = [p.text.strip() for p in doc.paragraphs]
    marker = "8. 重庆绕城高速交通流量预测与大屏可视化（9月9日）"
    if marker in existing:
        raise RuntimeError("报告中已存在 9 月 9 日任务章节，避免重复写入。")

    reference = next(
        (p for p in doc.paragraphs if p.text.strip() == "三、 实训总结与心得体会"),
        None,
    )
    if reference is None:
        raise RuntimeError("未找到“ 三、 实训总结与心得体会 ”章节。")

    add_text_before(doc, reference, marker, style="Heading 2")
    add_text_before(
        doc,
        reference,
        "• 任务关系分析与工程组织：分析重庆绕城高速交通流量预测模块与当前智能交通主项目的业务和技术边界。该模块使用 CQSkyEyeX 无人机航拍轨迹数据进行离线时序分析，主项目则面向视频流、车辆车牌识别和违章告警，两者在交通业务上互补，但数据形态、运行方式和依赖栈差异明显。因此将其作为独立实训工程运行，保留与主项目通过时序数据库或分析结果进行数据级对接的扩展空间。",
    )
    add_text_before(
        doc,
        reference,
        "• 轨迹数据读取与末端异常清洗：在 location1 目录中读取 1-1_trajectory.xlsx，原始数据共 336,166 条记录，字段包括 Time、ID、Velocity、LaneID 和 Class。检查时间戳后发现数据最大时间为 1200.72 秒，最后一个 60 秒时间窗实际只有 0.72 秒，属于采集结束造成的残缺窗口。程序按时间窗有效采样时长进行边界清洗，过滤小于 30 秒的残缺尾窗，最终保留 336,031 条有效记录和 20 个完整时间窗。",
    )
    add_text_before(
        doc,
        reference,
        "• 交通流量聚合与时间轴对齐：将轨迹数据按 60 秒切分时间窗，并在每个窗口内统计不同车辆 ID 的数量，作为该窗口的交通流量。为避免统计点与预测起点相差一个时间窗，绘图时将每个统计值定位到窗口结束时刻，使最后一个完整窗口 1140—1200 秒对应的历史流量点落在 1200 秒处。清洗后历史流量平均值为 56.2 辆/60 秒，最大值为 63 辆/60 秒，末端不再出现由残缺窗口引起的蓝线断崖。",
    )
    add_text_before(
        doc,
        reference,
        "• 预测模型与结果输出：采用 5 个时间窗的移动平均方法削弱离散轨迹计数的短时波动，再使用 Scikit-learn LinearRegression 对平滑后的历史序列进行趋势拟合，从 1200 秒开始向后预测 15 个时间步，即未来 15 分钟的交通流量变化。预测结果整体呈平缓下降趋势，未来 15 步平均流量约为 50.9 辆/60 秒，并通过非 GUI 模式输出预测大屏图片，同时保留 Tkinter 图形界面用于选择数据文件、调整时间窗、预测步数和移动平均窗口。",
    )
    add_text_before(
        doc,
        reference,
        "图 8-1 重庆绕城高速交通流量时序预测与可视化监控大屏：",
    )
    add_image_before(doc, reference, IMAGE)
    add_text_before(doc, reference, "")

    summary = next(
        (p for p in doc.paragraphs if p.text.strip().startswith("经过为期多天的实训")),
        None,
    )
    if summary is not None and "交通流量时序预测" not in summary.text:
        summary.add_run(
            "此外，本次完成的重庆绕城高速轨迹流量预测与大屏可视化实践，进一步补充了从交通感知数据到时序预测和决策展示的分析链路。"
        )

    doc.save(str(REPORT))
    print(f"updated: {REPORT}")
    print(f"paragraphs: {len(doc.paragraphs)}; inline_shapes: {len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
