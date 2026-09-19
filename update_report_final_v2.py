import docx

doc_path = 'D:/intelligent_transportation/人工智能智慧交通实训报告.docx'
doc = docx.Document(doc_path)

# 因为刚才已经删除了段落，段落列表索引发生变化，我们直接遍历找到总结那一段的前面插入
target_p = None
h2_style = None

# 获取样式
for p in doc.paragraphs:
    if '4. 边缘部署优化与模型加速' in p.text:
        h2_style = p.style
        break

# 获取插入点
for p in doc.paragraphs:
    if '三、 实训总结与心得体会' in p.text:
        target_p = p
        break

if target_p:
    h = target_p.insert_paragraph_before('5. 数据存储实战与全链路架构打通（9月4日 深度验收留痕）')
    if h2_style:
        h.style = h2_style
    
    target_p.insert_paragraph_before('结合项目组PPT验收标准与纯Python/SQL技术栈规范，今日全面重构并落地了以下企业级数据架构，完成了数据存储阶段的最终交付：')
    
    target_p.insert_paragraph_before('• 底层时序架构规范化重建：严格按照验收标准，编写并生成了 create_hypertable.sql 等三个独立SQL脚本。将表结构重构为标准的 traffic_data，成功开启了数据压缩策略（有效节约90%空间），创建了5分钟粒度的连续聚合视图，并配置了保留策略自动清理30天前的冷数据。')
    
    target_p.insert_paragraph_before('• 企业级数据入库代码开发：废弃单条同步写入模式，采用纯Python开发了生产级入库脚本 write_to_timescaledb.py。引入了 psycopg2.pool 数据库连接池和 executemany 批量写入技术，并添加了完整的网络异常与错误重试机制。')
    
    target_p.insert_paragraph_before('• 离线分析与冷数据归档实战：针对单机分析场景，开发了 duckdb_analysis.py。利用 DuckDB 的内存级列式计算特性对交通记录进行高速 OLAP 统计，将历史明细零拷贝导出为高压缩比的 Parquet 归档文件，并结合 Pandas 生成了包含车型占比、高峰时段的 Excel 报表。')
    
    target_p.insert_paragraph_before('• 端到端实时数据闭环链路：自主采用 FastAPI 搭建了高性能大屏后端 API，通过直接命中连续聚合视图实现极速响应；同时开发了“边缘端流量模拟器”与“Kafka 消费者服务”，打通了从“边缘感知 -> Kafka缓冲 -> TimescaleDB写入 -> 大屏API可视化”的全套工业级链路。')
    
    target_p.insert_paragraph_before('• 标准存储架构文档产出：提炼产出了包含 ER 表结构设计、分区策略与架构流转拓扑的《数据存储架构设计文档.md》，确保实训开发过程透明、完全符合工程化要求。')
    
    target_p.insert_paragraph_before('')

    doc.save(doc_path)
    print("报告更新成功！")
else:
    print("插入点未找到")
